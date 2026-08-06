import logging
import os
from typing import List, Dict, Any, Optional
import pdfplumber
import docx
import markdown
from bs4 import BeautifulSoup

logger = logging.getLogger(__name__)


class DocumentParser:
    def __init__(self, chunk_size: int = 512, chunk_overlap: int = 50):
        self.chunk_size = chunk_size
        self.chunk_overlap = chunk_overlap

    def parse_file(self, file_path: str) -> List[str]:
        _, ext = os.path.splitext(file_path)
        ext = ext.lower()

        try:
            if ext == ".pdf":
                return self._parse_pdf(file_path)
            elif ext == ".docx":
                return self._parse_docx(file_path)
            elif ext == ".md":
                return self._parse_markdown(file_path)
            elif ext == ".txt":
                return self._parse_txt(file_path)
            else:
                logger.warning(f"Unsupported file type: {ext}")
                return []
        except Exception as e:
            logger.error(f"Failed to parse {file_path}: {str(e)}")
            return []

    def _parse_pdf(self, file_path: str) -> List[str]:
        texts = []
        with pdfplumber.open(file_path) as pdf:
            for page in pdf.pages:
                text = page.extract_text()
                if text:
                    texts.append(text.strip())
                # 表格提取：Datasheet 的引脚表 / 参数表往往是核心检索对象
                try:
                    tables = page.extract_tables()
                    for tbl in tables:
                        if tbl:
                            texts.append(self._format_table(tbl))
                except Exception as e:
                    logger.warning(f"表格提取失败（{os.path.basename(file_path)}）: {e}")
        return texts

    @staticmethod
    def _format_table(table: List[List[Any]]) -> str:
        """将 pdfplumber 的表格（行列表）格式化为可检索文本。

        单元格以 ' | ' 分隔，行以换行分隔，并加 [TABLE] 标记，
        使引脚名 / 电气参数等结构化信息可被 RAG 检索到。
        """
        rows = []
        for row in table:
            cells = [str(c).replace("\n", " ").strip() if c is not None else "" for c in row]
            rows.append(" | ".join(cells))
        return "[TABLE]\n" + "\n".join(rows)

    def _parse_docx(self, file_path: str) -> List[str]:
        """解析 DOCX：正文合成整篇文本 + 表格单独抽取。

        正文不能逐段返回：一段一 chunk 会把「Supply voltage: 1.8 V to 3.6 V」这类
        关键参数行切成几个词的碎片，检索时排不进 top-k。表格（引脚表 / 电气参数表）
        与 PDF 一样用 [TABLE] 标记单独入库。
        """
        doc = docx.Document(file_path)
        paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        texts = ["\n".join(paragraphs)] if paragraphs else []
        for table in doc.tables:
            rows = [[cell.text for cell in row.cells] for row in table.rows]
            if rows:
                texts.append(self._format_table(rows))
        return texts

    def _parse_markdown(self, file_path: str) -> List[str]:
        with open(file_path, "r", encoding="utf-8") as f:
            content = f.read()
        
        md = markdown.Markdown(extensions=["extra"])
        html = md.convert(content)
        
        soup = BeautifulSoup(html, "html.parser")
        text = soup.get_text()
        
        return [text.strip()]

    def _parse_txt(self, file_path: str) -> List[str]:
        with open(file_path, "r", encoding="utf-8") as f:
            content = f.read()
        return [content.strip()]

    def split_text(self, text: str) -> List[str]:
        chunks = []
        start = 0
        text_length = len(text)

        while start < text_length:
            end = start + self.chunk_size
            if end >= text_length:
                chunk = text[start:]
                if chunk.strip():
                    chunks.append(chunk.strip())
                break

            last_period = text.rfind(". ", start, end)
            last_newline = text.rfind("\n", start, end)
            last_space = text.rfind(" ", start, end)

            split_pos = max(last_period, last_newline, last_space)
            if split_pos > start:
                end = split_pos + 1
            else:
                end = start + self.chunk_size

            chunk = text[start:end].strip()
            if chunk:
                chunks.append(chunk)

            start = end - self.chunk_overlap
            if start < 0:
                start = 0

        return chunks

    def process_file(
        self, file_path: str, metadata_extra: Optional[Dict[str, Any]] = None
    ) -> List[Dict[str, Any]]:
        raw_texts = self.parse_file(file_path)
        all_chunks = []

        for text in raw_texts:
            chunks = self.split_text(text)
            all_chunks.extend(chunks)

        documents = []
        for i, chunk in enumerate(all_chunks):
            metadata = {
                # source 与 file_name 同值：前端引用展示按 source 取，保留 file_name 兼容旧索引
                "source": os.path.basename(file_path),
                "file_name": os.path.basename(file_path),
                "chunk_index": i,
                "total_chunks": len(all_chunks),
            }
            if metadata_extra:
                metadata.update(metadata_extra)
            documents.append({
                "content": chunk,
                "metadata": metadata,
            })

        return documents